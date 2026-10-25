# src/auditoria_excel.py
"""
Excel profesional de la trazabilidad:

  · generar_excel_retiro()     registro de un retiro de archivos (resumen + archivo por archivo)
  · generar_excel_auditoria()  control de actividad (resumen, por persona, detalle, descargas, retiros)

Mismo estilo corporativo que el resumen de cumplimiento (src/resumen_excel.py).
Los textos que escribe una persona se guardan como texto (nunca como fórmula).
"""

import io
from datetime import datetime
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.auditoria import ETIQUETAS, CATEGORIAS, Auditoria, ahora_lima
from src.resumen_excel import (AMBAR_BG, AMBAR_FG, AZUL, CLARO, NAVY, ROJO_BG, ROJO_FG, VERDE_BG, VERDE_FG,
                               _celda, _encabezados)
from src.solicitudes_password import fecha_lima


def _titulo(ws, fila: int, n_col: int, texto: str, tam: int, color: str, fondo: Optional[str] = None, alto: Optional[int] = None):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_col)
    c = ws.cell(row=fila, column=1)
    c.value = texto
    c.data_type = "s"
    c.font = Font(name="Calibri", size=tam, bold=True, color=color)
    if fondo:
        for col in range(1, n_col + 1):
            ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=fondo)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if alto:
        ws.row_dimensions[fila].height = alto


def _hoja(wb: Workbook, titulo_hoja: str, anchos: List[int], primera: bool = False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo_hoja[:31]
    ws.sheet_view.showGridLines = False
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    return ws


def _cabecera(ws, n_col: int, titulo: str, sub: str):
    _titulo(ws, 1, n_col, "EsSalud · Oficina de Inteligencia e Información Sanitaria (OIIS)", 10, "FFFFFF", AZUL, 20)
    _titulo(ws, 2, n_col, titulo, 16, "FFFFFF", NAVY, 30)
    _titulo(ws, 3, n_col, sub, 10, NAVY, CLARO, 22)


def _pie(ws, fila: int, n_col: int, texto: str):
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_col)
    c = ws.cell(row=fila, column=1)
    c.value = texto
    c.data_type = "s"
    c.font = Font(name="Calibri", size=9, italic=True, color="566573")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[fila].height = 30


def _formato_impresion(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _pares(ws, fila: int, pares: List[tuple], col_valor_fin: int) -> int:
    for etiqueta, valor in pares:
        _celda(ws, fila, 1, etiqueta, negrita=True, fondo=CLARO)
        for col in range(2, col_valor_fin + 1):      # bordes de toda la celda combinada
            _celda(ws, fila, col, valor if col == 2 else None)
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=col_valor_fin)
        ws.row_dimensions[fila].height = max(18, 15 * (1 + len(str(valor)) // 90))
        fila += 1
    return fila


def _tono(valor: str):
    v = (valor or "").upper()
    if v in ("OK", "EJECUTADA", "DESCARGADA"):
        return VERDE_BG, VERDE_FG
    if v in ("ERROR", "RECHAZADO", "RECHAZADA"):
        return ROJO_BG, ROJO_FG
    return AMBAR_BG, AMBAR_FG


def nombre_excel_retiro(sol: Dict) -> str:
    return f"Registro_retiro_N{int(sol.get('id', 0)):04d}_S{int(sol.get('semana', 0)):02d}_{int(sol.get('anio', 0))}.xlsx"


def generar_excel_retiro(sol: Dict, archivos: List[Dict], resumen: Optional[Dict] = None) -> bytes:
    """Registro de un retiro: hoja «Resumen» y hoja «Archivos»."""
    resumen = resumen or {}
    wb = Workbook()
    n = 8
    ws = _hoja(wb, "Resumen del retiro", [26, 20, 20, 20, 20, 20, 20, 20], primera=True)
    _cabecera(ws, n, "Registro de retiro de archivos del repositorio",
              f"Solicitud N.º {sol.get('id')} · Semana {sol.get('semana')} · Año {sol.get('anio')} · "
              f"Generado el {ahora_lima():%d/%m/%Y %H:%M} (hora de Lima)")
    redes = "Todas las redes" if sol.get("todas_las_redes") else ", ".join(sol.get("redes") or []) or "—"
    n_ok = int(resumen.get("n_ok", sol.get("n_ok", 0)) or 0)
    n_err = int(resumen.get("n_error", sol.get("n_error", 0)) or 0)
    alertas = ", ".join(resumen.get("alertas") or []) or "Ninguna"
    fila = 5
    _titulo(ws, fila, n, "1. Datos de la solicitud", 12, NAVY, None, 22)
    fila = _pares(ws, fila + 1, [
        ("Solicitud N.º", str(sol.get("id"))),
        ("Semana y año", f"Semana {sol.get('semana')} · {sol.get('anio')}"),
        ("Redes", redes),
        ("Envíos incluidos", str(sol.get("n_envios", len(sol.get("envio_ids") or [])))),
        ("Motivo (categoría)", sol.get("motivo_categoria") or "—"),
        ("Motivo (detalle)", sol.get("motivo") or "—"),
        ("Descarga previa", "Sí, descargados antes del retiro" if not sol.get("sin_descarga_previa") else "NO: se retiraron sin descarga previa"),
        ("Tipo de borrado", "Definitivo (no recuperable)" if sol.get("definitivo") else "A la papelera de Drive (recuperable unos 30 días)"),
    ], n)
    fila += 1
    _titulo(ws, fila, n, "2. Responsables y fechas (hora de Lima)", 12, NAVY, None, 22)
    fila = _pares(ws, fila + 1, [
        ("Solicitó", f"{sol.get('solicitado_nombre') or sol.get('solicitado_por')} ({sol.get('solicitado_por')}) · {sol.get('rol_solicitante') or ''}"),
        ("Fecha de la solicitud", fecha_lima(sol.get("fecha_creacion"))),
        ("Aprobó / ejecutó", f"{sol.get('decidido_nombre') or sol.get('decidido_por') or '—'} ({sol.get('decidido_por') or '—'})"),
        ("Fecha de la decisión", fecha_lima(sol.get("fecha_decision"))),
        ("Fecha de ejecución", fecha_lima(sol.get("fecha_ejecucion"))),
        ("Nota de la decisión", sol.get("nota_decision") or "—"),
        ("Equipo (IP)", sol.get("ip") or "no disponible"),
    ], n)
    fila += 1
    _titulo(ws, fila, n, "3. Resultado", 12, NAVY, None, 22)
    fila = _pares(ws, fila + 1, [
        ("Estado", sol.get("estado") or "—"),
        ("Archivos retirados", str(n_ok)),
        ("Archivos con error", str(n_err)),
        ("Señales para revisar", alertas),
    ], n)
    fila += 1
    _pie(ws, fila, n, "Se conserva el registro de recepción (quién envió y cuándo) en la base de datos; solo se quitó el archivo del repositorio. "
                      "Este documento es un comprobante generado por SIREA; el original está en la auditoría del sistema.")
    _formato_impresion(ws)

    ws2 = _hoja(wb, "Archivos", [6, 22, 22, 9, 8, 38, 34, 26, 12, 40, 18])
    n2 = 11
    _cabecera(ws2, n2, "Detalle archivo por archivo", f"Solicitud N.º {sol.get('id')} · Semana {sol.get('semana')} · Año {sol.get('anio')}")
    fila = 5
    _encabezados(ws2, fila, ["N.º", "Código de seguimiento", "Red", "Semana", "Año", "Centros", "Ruta del ZIP", "ID en Drive",
                             "Resultado", "Detalle", "Fecha y hora"])
    ws2.freeze_panes = ws2.cell(row=fila + 1, column=3)
    for i, a in enumerate(archivos, 1):
        fila += 1
        zebra = CLARO if i % 2 == 0 else None
        _celda(ws2, fila, 1, i, alinea="center", fondo=zebra)
        _celda(ws2, fila, 2, a.get("codigo_seguimiento") or "—", negrita=True, fondo=zebra)
        _celda(ws2, fila, 3, a.get("red") or "—", fondo=zebra)
        _celda(ws2, fila, 4, a.get("semana"), alinea="center", fondo=zebra)
        _celda(ws2, fila, 5, a.get("anio"), alinea="center", fondo=zebra)
        _celda(ws2, fila, 6, a.get("centros") or "—", fondo=zebra)
        _celda(ws2, fila, 7, a.get("ruta_zip") or "—", fondo=zebra)
        _celda(ws2, fila, 8, a.get("drive_file_id") or "—", fondo=zebra)
        bg, fg = _tono(a.get("resultado"))
        _celda(ws2, fila, 9, "Retirado" if a.get("resultado") == "OK" else "Error", negrita=True, color=fg, fondo=bg, alinea="center")
        _celda(ws2, fila, 10, a.get("detalle") or "", fondo=zebra)
        _celda(ws2, fila, 11, fecha_lima(a.get("fecha")), alinea="center", fondo=zebra)
    if not archivos:
        fila += 1
        ws2.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n2)
        _celda(ws2, fila, 1, "Aún no hay archivos registrados.", alinea="center")
    _formato_impresion(ws2)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _txt_detalle(f: Dict) -> str:
    d = Auditoria.detalle_de(f)
    partes = []
    for k, v in d.items():
        if k in ("alertas", "codigos", "items"):
            continue
        if isinstance(v, (list, tuple)):
            v = ", ".join(str(x) for x in v)
        partes.append(f"{k}: {v}")
    return " · ".join(partes)[:600]


def generar_excel_auditoria(filas: List[Dict], filtros: Dict, por_persona: List[Dict], descargas: List[Dict],
                            retiros: List[Dict], cadena: Optional[Dict] = None, generado_por: str = "") -> bytes:
    """Control de actividad. `filas` = filas de auditoría (más recientes primero)."""
    wb = Workbook()
    n = 7
    ws = _hoja(wb, "Resumen", [30, 18, 18, 18, 18, 18, 18], primera=True)
    _cabecera(ws, n, "Control de actividad del panel", f"Generado el {ahora_lima():%d/%m/%Y %H:%M} (hora de Lima)"
              + (f" por {generado_por}" if generado_por else ""))
    con_alerta = [f for f in filas if Auditoria.detalle_de(f).get("alertas")]
    fila = 5
    _titulo(ws, fila, n, "1. Filtros aplicados", 12, NAVY, None, 22)
    fila = _pares(ws, fila + 1, [(k, str(v) if v else "Todos") for k, v in filtros.items()], n)
    fila += 1
    _titulo(ws, fila, n, "2. Indicadores", 12, NAVY, None, 22)
    fila = _pares(ws, fila + 1, [
        ("Acciones registradas", str(len(filas))),
        ("Personas distintas", str(len({f.get('usuario') for f in filas}))),
        ("Descargas masivas", str(sum(1 for f in filas if f.get("accion") == "DESCARGA_MASIVA"))),
        ("Retiros ejecutados", str(sum(1 for f in filas if f.get("accion") == "RETIRO_EJECUTADO"))),
        ("Acciones con señales para revisar", str(len(con_alerta))),
        ("Intentos sin permiso", str(sum(1 for f in filas if f.get("accion") == "ACCESO_DENEGADO"))),
    ], n)
    fila += 1
    _titulo(ws, fila, n, "3. Integridad del registro", 12, NAVY, None, 22)
    if cadena is None:
        estado = "No se verificó al generar este archivo."
    elif cadena.get("ok"):
        estado = f"ÍNTEGRO: la cadena de {cadena.get('filas')} registros no muestra alteraciones."
    else:
        estado = f"ALTERADO: {cadena.get('motivo')} (registro N.º {cadena.get('id_roto')})."
    fila = _pares(ws, fila + 1, [("Resultado", estado)], n)
    fila += 1
    _pie(ws, fila, n, "La auditoría es de solo escritura y cada registro lleva la huella del anterior. Los acuses por correo son un registro independiente.")
    _formato_impresion(ws)

    ws1 = _hoja(wb, "Por persona", [24, 16, 12, 12, 12, 12, 14, 20])
    _cabecera(ws1, 8, "Acciones por persona", "Ordenado por señales para revisar y cantidad de acciones")
    _encabezados(ws1, 5, ["Usuario", "Rol", "Acciones", "Descargas", "Retiros", "Duplicados", "Señales", "Última acción"])
    r = 5
    for i, p in enumerate(por_persona, 1):
        r += 1
        z = CLARO if i % 2 == 0 else None
        _celda(ws1, r, 1, p["usuario"], negrita=True, fondo=z)
        _celda(ws1, r, 2, p.get("rol") or "—", alinea="center", fondo=z)
        for c, k in enumerate(("total", "descargas", "retiros", "duplicados"), 3):
            _celda(ws1, r, c, p[k], alinea="center", fondo=z)
        bg, fg = (ROJO_BG, ROJO_FG) if p["alertas"] else (VERDE_BG, VERDE_FG)
        _celda(ws1, r, 7, p["alertas"], negrita=True, color=fg, fondo=bg, alinea="center")
        _celda(ws1, r, 8, _fmt_iso(p.get("ultima")), alinea="center", fondo=z)
    ws1.freeze_panes = "A6"
    _formato_impresion(ws1)

    ws2 = _hoja(wb, "Detalle de acciones", [6, 19, 20, 11, 22, 26, 38, 11, 28, 40, 16])
    _cabecera(ws2, 11, "Detalle de acciones", f"{len(filas)} registro(s), del más reciente al más antiguo")
    _encabezados(ws2, 5, ["N.º", "Fecha y hora (Lima)", "Usuario", "Rol", "Categoría", "Acción", "Objeto", "Resultado",
                          "Señales", "Detalle", "IP"])
    r = 5
    for f in filas:
        r += 1
        z = CLARO if r % 2 == 0 else None
        d = Auditoria.detalle_de(f)
        _celda(ws2, r, 1, f.get("id"), alinea="center", fondo=z)
        _celda(ws2, r, 2, _fmt_iso(f.get("ocurrio_iso")), alinea="center", fondo=z)
        _celda(ws2, r, 3, f.get("usuario"), negrita=True, fondo=z)
        _celda(ws2, r, 4, f.get("rol") or "—", alinea="center", fondo=z)
        _celda(ws2, r, 5, CATEGORIAS.get(f.get("categoria"), f.get("categoria")), fondo=z)
        _celda(ws2, r, 6, ETIQUETAS.get(f.get("accion"), f.get("accion")), fondo=z)
        _celda(ws2, r, 7, f.get("objeto") or "", fondo=z)
        bg, fg = _tono(f.get("resultado"))
        _celda(ws2, r, 8, f.get("resultado"), negrita=True, color=fg, fondo=bg, alinea="center")
        al = ", ".join(d.get("alertas") or [])
        _celda(ws2, r, 9, al, color=ROJO_FG if al else None, negrita=bool(al), fondo=z)
        _celda(ws2, r, 10, _txt_detalle(f), fondo=z)
        _celda(ws2, r, 11, f.get("ip") or "", fondo=z)
    ws2.freeze_panes = "C6"
    _formato_impresion(ws2)

    ws3 = _hoja(wb, "Descargas masivas", [6, 19, 22, 8, 8, 28, 40, 9, 11, 12, 66])
    _cabecera(ws3, 11, "Descargas masivas", "ZIP preparados y descargados")
    _encabezados(ws3, 5, ["N.º", "Preparada", "Responsable", "Semana", "Año", "Redes", "Archivo ZIP", "Archivos", "Estado", "Tamaño (MB)", "Huella SHA-256"])
    r = 5
    for d in descargas:
        r += 1
        z = CLARO if r % 2 == 0 else None
        _celda(ws3, r, 1, d.get("id"), alinea="center", fondo=z)
        _celda(ws3, r, 2, fecha_lima(d.get("fecha_preparada")), alinea="center", fondo=z)
        _celda(ws3, r, 3, d.get("usuario"), negrita=True, fondo=z)
        _celda(ws3, r, 4, d.get("semana"), alinea="center", fondo=z)
        _celda(ws3, r, 5, d.get("anio"), alinea="center", fondo=z)
        _celda(ws3, r, 6, "Todas" if d.get("todas_las_redes") else ", ".join(d.get("redes") or []), fondo=z)
        _celda(ws3, r, 7, d.get("nombre_zip"), fondo=z)
        _celda(ws3, r, 8, d.get("n_archivos"), alinea="center", fondo=z)
        bg, fg = _tono(d.get("estado"))
        _celda(ws3, r, 9, d.get("estado"), negrita=True, color=fg, fondo=bg, alinea="center")
        _celda(ws3, r, 10, round(float(d.get("bytes") or 0) / 1024 ** 2, 1), alinea="center", fondo=z, formato="0.0")
        _celda(ws3, r, 11, d.get("sha256") or "", fondo=z)
    ws3.freeze_panes = "C6"
    _formato_impresion(ws3)

    ws4 = _hoja(wb, "Retiros", [6, 19, 22, 22, 8, 8, 12, 10, 10, 30, 54])
    _cabecera(ws4, 11, "Retiros de archivos del repositorio", "Solicitudes y su resultado")
    _encabezados(ws4, 5, ["N.º", "Solicitada", "Solicitó", "Aprobó / ejecutó", "Semana", "Año", "Estado", "Retirados", "Con error", "Motivo", "Detalle del motivo"])
    r = 5
    for s in retiros:
        r += 1
        z = CLARO if r % 2 == 0 else None
        _celda(ws4, r, 1, s.get("id"), alinea="center", fondo=z)
        _celda(ws4, r, 2, fecha_lima(s.get("fecha_creacion")), alinea="center", fondo=z)
        _celda(ws4, r, 3, s.get("solicitado_por"), negrita=True, fondo=z)
        _celda(ws4, r, 4, s.get("decidido_por") or "—", fondo=z)
        _celda(ws4, r, 5, s.get("semana"), alinea="center", fondo=z)
        _celda(ws4, r, 6, s.get("anio"), alinea="center", fondo=z)
        bg, fg = _tono(s.get("estado"))
        _celda(ws4, r, 7, s.get("estado"), negrita=True, color=fg, fondo=bg, alinea="center")
        _celda(ws4, r, 8, s.get("n_ok"), alinea="center", fondo=z)
        _celda(ws4, r, 9, s.get("n_error"), alinea="center", fondo=z)
        _celda(ws4, r, 10, s.get("motivo_categoria"), fondo=z)
        _celda(ws4, r, 11, s.get("motivo"), fondo=z)
    ws4.freeze_panes = "C6"
    _formato_impresion(ws4)

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _fmt_iso(valor) -> str:
    """'2026-10-15 09:30:00' -> '15/10/2026 09:30:00'."""
    try:
        return datetime.strptime(str(valor), "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y %H:%M:%S")
    except (ValueError, TypeError):
        return str(valor or "—")


def nombre_excel_auditoria() -> str:
    return f"Control_de_actividad_{ahora_lima():%Y%m%d_%H%M}.xlsx"
