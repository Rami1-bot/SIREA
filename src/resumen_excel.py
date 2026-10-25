# src/resumen_excel.py
"""
Excel «Resumen de cumplimiento por red».

Una hoja con dos tablas:

  1. Una fila por red de la semana pedida: centros esperados y con envío,
     porcentaje, estado, quiénes hicieron los envíos y cuándo completaron
     (o cuándo fue el último envío si aún falta algún centro).
  2. Debajo, los envíos TARDÍOS de semanas anteriores que se registraron
     mientras esa semana estaba activa (Fecha y hora, red, centros, semana a la
     que corresponden y días de atraso).

La recolección de datos (`recolectar`) está separada del dibujo del Excel
(`generar_excel`) para poder probarlas por separado y mostrar una vista
previa en el panel.
"""

import io
from datetime import datetime, timedelta, timezone
from typing import Callable, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.calendario_epi import formatear_rango, rango_semana
from src.solicitudes_password import fecha_lima, parsear_fecha
from utils.duplicados import solo_originales

_LIMA = timezone(timedelta(hours=-5))

NAVY = "1D3F76"
AZUL = "006DB0"
CLARO = "F4F6F8"
LINEA = "D5DDE4"
VERDE_BG, VERDE_FG = "E7F3EA", "1E6B3A"
AMBAR_BG, AMBAR_FG = "FBF3E4", "7A5A00"
ROJO_BG, ROJO_FG = "FDECEA", "B42318"


# ------------------------------------------------------------------ datos

def _nombre_persona(envio: Dict) -> str:
    return f"{envio.get('nombre_usuario', '') or ''} {envio.get('apellido_usuario', '') or ''}".strip()


def _etiqueta_centros(envio: Dict) -> str:
    centros = envio.get("centros") or []
    if not isinstance(centros, list):
        return str(centros)
    partes = []
    for c in centros:
        if isinstance(c, dict):
            partes.append(f"{c.get('codigo', '')} {c.get('nombre', '')}".strip())
        else:
            partes.append(str(c))
    return ", ".join(partes)


def _nombres_limpios(envio: Dict, limpiar: Callable[[str], str]) -> List[str]:
    nombres = []
    for c in envio.get("centros") or []:
        n = c.get("nombre", str(c)) if isinstance(c, dict) else limpiar(str(c))
        if n:
            nombres.append(n.strip().lower())
    return nombres


def _fecha(valor) -> Optional[datetime]:
    d = parsear_fecha(valor)
    return d.astimezone(_LIMA) if d else None


def _txt(d: Optional[datetime]) -> str:
    return d.strftime("%d/%m/%Y %H:%M:%S") if d else "—"


def recolectar(db, anio: int, semana: int, redes: List[str], corte: Optional[datetime] = None,
               avance_fn: Optional[Callable] = None, limpiar_fn: Optional[Callable] = None) -> Dict:
    """
    Reúne los datos del resumen. `corte` (hora de Lima, sin zona) permite marcar
    los envíos hechos después del corte. `avance_fn(red, anio, semana)` devuelve
    {'total','enviados','pendientes'} como el resto del panel; por defecto se usa
    el mismo cálculo del sistema.
    """
    if avance_fn is None or limpiar_fn is None:
        from app_pages.usuario_logica import calcular_avance_red, extraer_nombre_limpio
        avance_fn = avance_fn or calcular_avance_red
        limpiar_fn = limpiar_fn or extraer_nombre_limpio

    filas = []
    for red in sorted(redes):
        av = avance_fn(red, anio, semana)
        if not av or not av.get("total"):
            continue
        total, enviados = int(av["total"]), list(av.get("enviados", []))
        envios = solo_originales(db.obtener_envios(anio=anio, semana=semana, red=red, limit=2000) or [])
        envios = sorted(envios, key=lambda e: str(e.get("fecha_envio") or ""))

        # Primer envío que cubrió cada centro (por nombre limpio, sin distinguir mayúsculas).
        primero: Dict[str, datetime] = {}
        for e in envios:
            f = _fecha(e.get("fecha_envio"))
            for n in _nombres_limpios(e, limpiar_fn):
                if f and n not in primero:
                    primero[n] = f
        nombres_enviados = {limpiar_fn(x).strip().lower() for x in enviados}
        tiempos = [primero[n] for n in nombres_enviados if n in primero]

        n_env = len(enviados)
        if n_env >= total:
            estado = "Completo"
            cierre = max(tiempos) if tiempos else None
        elif n_env > 0:
            estado = "En proceso"
            cierre = max((_fecha(e.get("fecha_envio")) for e in envios if _fecha(e.get("fecha_envio"))), default=None)
        else:
            estado, cierre = "Sin envíos", None

        responsables = []
        for e in envios:
            p = _nombre_persona(e)
            if p and p not in responsables:
                responsables.append(p)

        despues = bool(corte and cierre and cierre.replace(tzinfo=None) > corte)
        filas.append({
            "red": red, "esperados": total, "enviados": n_env,
            "pct": round(n_env * 100 / total, 1) if total else 0.0,
            "estado": estado, "responsables": ", ".join(responsables) or "—",
            "cierre": cierre, "cierre_etiqueta": "Completó" if estado == "Completo" else "Último envío",
            "despues_corte": despues, "n_envios": len(envios),
        })

    tardios = []
    for e in db.obtener_envios_atrasados(semana, limit=2000, anio=anio) or []:
        tardios.append({
            "red": e.get("red", ""), "centros": _etiqueta_centros(e), "semana_reporte": e.get("semana"),
            "responsable": _nombre_persona(e) or "—", "fecha": _fecha(e.get("fecha_envio")),
            "dias_atraso": e.get("dias_atraso_envio") or 0, "codigo": e.get("codigo_seguimiento", ""),
        })
    tardios.sort(key=lambda t: (t["red"], t["fecha"] or datetime.min.replace(tzinfo=_LIMA)))

    periodo = "—"
    try:
        info = db.obtener_anio_epidemiologico(int(anio))
        d, h = rango_semana(str(info["fecha_inicio"])[:10], semana)
        periodo = formatear_rango(d, h)
    except Exception:
        pass

    tot_esp = sum(f["esperados"] for f in filas)
    tot_env = sum(f["enviados"] for f in filas)
    return {
        "anio": anio, "semana": semana, "periodo": periodo, "corte": corte,
        "filas": filas, "tardios": tardios,
        "total_esperados": tot_esp, "total_enviados": tot_env,
        "pct_total": round(tot_env * 100 / tot_esp, 1) if tot_esp else 0.0,
        "generado": datetime.now(_LIMA),
    }


# ------------------------------------------------------------------ Excel

_borde = Border(*(Side(style="thin", color=LINEA) for _ in range(4)))


def _celda(ws, fila: int, col: int, valor, *, negrita=False, color=None, fondo=None, alinea="left",
           formato=None, ajustar=True):
    c = ws.cell(row=fila, column=col)
    c.value = valor
    # Un texto escrito por una persona que empiece con "=" no debe ejecutarse como fórmula.
    if isinstance(valor, str):
        c.data_type = "s"
    c.font = Font(name="Calibri", size=10, bold=negrita, color=color or "14202B")
    if fondo:
        c.fill = PatternFill("solid", fgColor=fondo)
    c.alignment = Alignment(horizontal=alinea, vertical="center", wrap_text=ajustar)
    c.border = _borde
    if formato:
        c.number_format = formato
    return c


def _encabezados(ws, fila: int, titulos: List[str]):
    for i, t in enumerate(titulos, 1):
        _celda(ws, fila, i, t, negrita=True, color="FFFFFF", fondo=NAVY, alinea="center")
    ws.row_dimensions[fila].height = 30


def _tono_pct(p: float):
    if p >= 90:
        return VERDE_BG, VERDE_FG
    if p >= 60:
        return AMBAR_BG, AMBAR_FG
    return ROJO_BG, ROJO_FG


def generar_excel(datos: Dict) -> bytes:
    """Dibuja el Excel corporativo a partir de `recolectar()`."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    anchos = [30, 11, 11, 13, 14, 44, 20, 22]
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    n_col = len(anchos)

    def titulo(fila, texto, tam, color, fondo=None, alto=None):
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_col)
        c = ws.cell(row=fila, column=1)
        c.value = texto
        c.data_type = "s"
        c.font = Font(name="Calibri", size=tam, bold=True, color=color)
        if fondo:
            for col in range(1, n_col + 1):
                ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=fondo)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        if alto:
            ws.row_dimensions[fila].height = alto

    titulo(1, "EsSalud · Oficina de Inteligencia e Información Sanitaria (OIIS)", 10, "FFFFFF", AZUL, 20)
    titulo(2, "Resumen de cumplimiento por red", 16, "FFFFFF", NAVY, 30)
    corte = datos.get("corte")
    sub = f"Semana {datos['semana']} · Año {datos['anio']} · Periodo {datos['periodo']}"
    if corte:
        sub += f" · Corte: {corte:%d/%m/%Y %H:%M:%S}"
    titulo(3, sub, 11, NAVY, CLARO, 22)
    titulo(4, f"Generado el {datos['generado']:%d/%m/%Y %H:%M} (hora de Lima) · "
              f"Cumplimiento nacional: {datos['pct_total']:.1f} % ({datos['total_enviados']} de {datos['total_esperados']} centros)",
           10, "566573", None, 18)

    # ---- Tabla 1
    fila = 6
    titulo(fila, "1. Cumplimiento de la semana por red", 12, NAVY, None, 22)
    fila += 1
    _encabezados(ws, fila, ["Red asistencial", "Centros esperados", "Centros con envío", "Cumplimiento",
                            "Estado", "Responsables del envío", "Completó / Último envío", "Observación"])
    ws.freeze_panes = ws.cell(row=fila + 1, column=2)
    for i, f in enumerate(datos["filas"]):
        fila += 1
        zebra = CLARO if i % 2 else None
        _celda(ws, fila, 1, f["red"], negrita=True, fondo=zebra)
        _celda(ws, fila, 2, f["esperados"], alinea="center", fondo=zebra)
        _celda(ws, fila, 3, f["enviados"], alinea="center", fondo=zebra)
        bg, fg = _tono_pct(f["pct"])
        _celda(ws, fila, 4, f["pct"] / 100, negrita=True, color=fg, fondo=bg, alinea="center", formato="0.0%")
        _celda(ws, fila, 5, f["estado"], alinea="center", fondo=zebra)
        _celda(ws, fila, 6, f["responsables"], fondo=zebra)
        if f["cierre"]:
            _celda(ws, fila, 7, f"{f['cierre_etiqueta']}: {_txt(f['cierre'])}", alinea="center", fondo=zebra)
        else:
            _celda(ws, fila, 7, "—", alinea="center", fondo=zebra)
        obs = "Posterior al corte" if f["despues_corte"] else ""
        _celda(ws, fila, 8, obs, color=ROJO_FG if obs else None, alinea="center", fondo=zebra)
    if not datos["filas"]:
        fila += 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_col)
        _celda(ws, fila, 1, "No hay redes con centros esperados para esta semana.", alinea="center")
    else:
        fila += 1
        _celda(ws, fila, 1, "TOTAL NACIONAL", negrita=True, color="FFFFFF", fondo=NAVY)
        _celda(ws, fila, 2, datos["total_esperados"], negrita=True, color="FFFFFF", fondo=NAVY, alinea="center")
        _celda(ws, fila, 3, datos["total_enviados"], negrita=True, color="FFFFFF", fondo=NAVY, alinea="center")
        _celda(ws, fila, 4, datos["pct_total"] / 100, negrita=True, color="FFFFFF", fondo=NAVY, alinea="center", formato="0.0%")
        for col in range(5, n_col + 1):
            _celda(ws, fila, col, "", fondo=NAVY)

    # ---- Tabla 2
    fila += 3
    titulo(fila, f"2. Envíos tardíos de semanas anteriores recibidos durante la semana {datos['semana']}", 12, NAVY, None, 22)
    fila += 1
    _encabezados(ws, fila, ["Red asistencial", "Semana del reporte", "Días de atraso", "Código de seguimiento",
                            "Fecha y hora de envío", "Centros", "Responsable del envío", ""])
    ws.cell(row=fila, column=8).fill = PatternFill(fill_type=None)
    ws.cell(row=fila, column=8).border = Border()
    if datos["tardios"]:
        for i, t in enumerate(datos["tardios"]):
            fila += 1
            zebra = CLARO if i % 2 else None
            _celda(ws, fila, 1, t["red"], negrita=True, fondo=zebra)
            _celda(ws, fila, 2, t["semana_reporte"], alinea="center", fondo=zebra)
            _celda(ws, fila, 3, t["dias_atraso"], alinea="center", fondo=zebra)
            _celda(ws, fila, 4, t["codigo"], alinea="center", fondo=zebra)
            _celda(ws, fila, 5, _txt(t["fecha"]), alinea="center", fondo=zebra)
            _celda(ws, fila, 6, t["centros"], fondo=zebra)
            _celda(ws, fila, 7, t["responsable"], fondo=zebra)
    else:
        fila += 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
        _celda(ws, fila, 1, "No se registraron envíos tardíos de semanas anteriores durante esta semana.", alinea="center")

    fila += 2
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_col)
    nota = ws.cell(row=fila, column=1)
    nota.value = ("Fuente: SIREA · Los porcentajes no incluyen los centros excluidos por casos excepcionales. "
                  "Los envíos duplicados no se cuentan como cumplimiento.")
    nota.font = Font(name="Calibri", size=9, italic=True, color="566573")
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[fila].height = 28

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def nombre_archivo(anio: int, semana: int, borrador: bool = False) -> str:
    return f"Resumen_cumplimiento_por_red_S{int(semana):02d}_{int(anio)}{'_borrador' if borrador else ''}.xlsx"
